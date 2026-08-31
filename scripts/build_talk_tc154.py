#!/usr/bin/env python3
"""Build the CC 36010 showcase deck for ISO/TC 154.

Uses the signatif TC 154 deck as the template: loads its presentation
(masters, layouts, theme), removes the original slides, and builds the
CC 36010 slides on the template's own layouts so backgrounds, title
placeholders, and fonts are the template's, not imitations.
"""
import copy
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = Path("/Users/mulgogi/src/calconnect/cc-signatif/presentations/"
                "20260831-calconnect-iso-tc-154-signatif.pptx")
DIAG = ROOT / "presentations" / "diagrams"

FONT = "Avenir Book"
INK = RGBColor(0x1B, 0x48, 0x69)
BLUE = RGBColor(0x25, 0x61, 0x8C)
BODY = RGBColor(0x4E, 0x5D, 0x6D)
MUTE = RGBColor(0x62, 0x6D, 0x79)
ACCENT = RGBColor(0xC7, 0x4A, 0x35)
CARD_LINE = RGBColor(0xC7, 0xD3, 0xDD)
PANEL = RGBColor(0xEA, 0xF1, 0xF6)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation(str(TEMPLATE))

# drop the template's own slides; keep masters, layouts, theme
for sldId in list(prs.slides._sldIdLst):
    prs.slides._sldIdLst.remove(sldId)

LAYOUTS = {}
for master in prs.slide_masters:
    for layout in master.slide_layouts:
        LAYOUTS.setdefault(layout.name, layout)

TITLE_LAYOUT = LAYOUTS["1_Section Header"]
CONTENT_LAYOUT = LAYOUTS["Title and Content"]
CLOSING_LAYOUT = LAYOUTS["Custom Layout"]

PAGE = [0]


def slide(layout):
    s = prs.slides.add_slide(layout)
    PAGE[0] += 1
    return s


def box(s, x, y, w, h, text, size=10.5, color=BODY, bold=False,
        align=None, italic=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for line in text.split("\n"):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = line
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.bold = bold
        p.font.italic = italic
        if align:
            p.alignment = align
    return tb


def rect(s, x, y, w, h, fill=PAPER, line=CARD_LINE, rounded=True):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        r.adjustments[0] = 0.06
    if fill is None:
        r.fill.background()
    else:
        r.fill.solid()
        r.fill.fore_color.rgb = fill
    if line is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = line
        r.line.width = Pt(0.75)
    return r


def bar(s, x, y, w=0.4, color=ACCENT):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                           Inches(w), Pt(2.4))
    r.fill.solid()
    r.fill.fore_color.rgb = color
    r.line.fill.background()
    return r


def set_title(s, text, lead=None):
    ph = s.placeholders[0]
    ph.text_frame.text = text
    if lead:
        box(s, 0.7, 0.95, 8.6, 0.55, lead, size=12.5, color=BODY)


def pageno(s):
    box(s, 9.3, 5.32, 0.3, 0.2, f"{PAGE[0]:02d}", size=6.5, color=MUTE)


def card(s, x, y, w, h, head, body, head_color=BLUE):
    rect(s, x, y, w, h)
    box(s, x + 0.12, y + 0.10, w - 0.24, 0.30, head,
        size=10.5, color=head_color, bold=True)
    bar(s, x + 0.13, y + 0.44)
    box(s, x + 0.12, y + 0.52, w - 0.24, h - 0.62, body, size=9.5)


def row(s, y, label, body, h=0.80):
    rect(s, 0.7, y, 8.6, h, fill=PANEL)
    box(s, 0.9, y + 0.10, 2.4, 0.5, label, size=10.5, color=INK, bold=True)
    box(s, 3.4, y + 0.07, 5.75, h - 0.12, body, size=9.5)


def diagram(s, name, x, y, w):
    s.shapes.add_picture(str(DIAG / name), Inches(x), Inches(y),
                         width=Inches(w))


# ---------------- 1 · Title (template section header) ----------------
s = slide(TITLE_LAYOUT)
s.placeholders[0].text_frame.text = "CC 36010: Lightweight document — Document metamodel"
sub = s.placeholders[1].text_frame
sub.text = "ISO/TC 154  ·  CalConnect"
sub.add_paragraph().text = "Published by CalConnect · fast-track showcase"
sub.add_paragraph().text = "Ronald TseAug 31, 2026"
for p in sub.paragraphs:
    for r in p.runs:
        r.font.name = FONT

# ---------------- 2 · The gap ----------------
s = slide(CONTENT_LAYOUT)
set_title(s, "Commerce exchanges everything except the document's text",
          "Every framework standardizes its data — none standardizes the rich text that carries the obligations.")
for i, (h, b) in enumerate([
    ("TRADE DATA", "EDIFACT, UBL, e-invoicing: structured fields, standardized — rich text rides as escaped HTML blobs."),
    ("PUBLISHING FORMATS", "HTML, RTF, DOCX: what documents look like — presentation contracts, not shared semantics."),
    ("MARKUP SYNTAXES", "Markdown, AsciiDoc, RST: convenient to type — no common meaning between them."),
    ("ELEMENT STANDARDS", "ISO 8601 time, ISO 639/15924/3166 language, ISO 690 citations — the pieces exist; the carrier does not."),
]):
    card(s, 0.7 + i * 2.2, 1.75, 2.05, 1.85, h, b)
rect(s, 0.7, 3.85, 8.6, 0.75, fill=PANEL)
box(s, 0.9, 3.98, 8.2, 0.55,
    "Exchange degrades to plain text: emphasis, terms, notes, tables and change intent lost on every hop.",
    size=13, color=INK, bold=True)
pageno(s)

# ---------------- 3 · What CC 36010 is ----------------
s = slide(CONTENT_LAYOUT)
set_title(s, "The missing layer: a metamodel of what documents mean",
          "Complete structure and semantics; presentation excluded by design and relegated to rendering.")
tiers = [
    ("Document", "identifier · bibdata (ISO 690 lineage) · attribute register", INK),
    ("Sections", "hierarchical · prefatory · references", BLUE),
    ("Blocks", "paragraphs · tables · lists · figures · source · admonitions · amend", ACCENT),
    ("Inline elements", "emphasis · stem · ruby · media · references · variables", RGBColor(0x2A, 0x6B, 0x7C)),
]
y = 1.7
for name, body, col in tiers:
    rect(s, 1.6, y, 6.8, 0.68)
    box(s, 1.8, y + 0.06, 2.1, 0.5, name, size=13, color=col, bold=True)
    box(s, 3.9, y + 0.10, 4.4, 0.5, body, size=9.5)
    y += 0.78
box(s, 0.7, 4.85, 8.6, 0.6,
    "One strict rule — nesting within a tier, never across tiers — is why every markup maps cleanly and rendering is predictable.",
    size=12, color=INK, bold=True)
pageno(s)

# ---------------- 4 · Serialization-agnostic ----------------
s = slide(CONTENT_LAYOUT)
set_title(s, "Model-driven: any serialization, zero-loss transforms",
          "The model binds to no format. XML and YAML are reference serializations; twins of identical content prove interchange.")
diagram(s, "d3-not-a-serialization.png", 1.25, 1.55, 7.5)
rect(s, 0.7, 4.62, 8.6, 0.62, fill=PANEL)
box(s, 0.9, 4.72, 8.2, 0.45,
    "Conformance is to the model — any language, any format; round-trip equivalence to a reference serialization is the test.",
    size=12, color=INK, bold=True)
pageno(s)

# ---------------- 5 · One model, many markups ----------------
s = slide(CONTENT_LAYOUT)
set_title(s, "Harmonized coverage, clause by clause (Annex A)",
          "Normative mapping tables: every construct of the principal markups has a home.")
for i, (h, b) in enumerate([
    ("ASCIIDOC", "admonitions → Admonition+type\ncontinuations → relaxed list items\npassthroughs → FormattedString\n{attr} → variable reference"),
    ("MARKDOWN · GFM · PANDOC", "GFM tables → TableBlock\n{#id .class} → attribute register\nfenced divs → register + container\nraw HTML → format-qualified strings"),
    ("RST · SPHINX", "directives → register + relaxed content\nsubstitutions → variable reference\nfield lists → document register\ndanger/error → type specializations"),
]):
    card(s, 0.7 + i * 2.95, 1.7, 2.75, 2.3, h, b)
rect(s, 0.7, 4.25, 8.6, 0.68, fill=PANEL)
box(s, 0.9, 4.36, 8.2, 0.5,
    "Unlisted constructs are covered by the conformance formula: construct · type-specialization · register + relaxed or opaque content.",
    size=12, color=INK, bold=True)
pageno(s)

# ---------------- 6 · Adopt three ways ----------------
s = slide(CONTENT_LAYOUT)
set_title(s, "Communities adopt it three ways — none touch the base",
          "Open-closed by construction: the published basis is never modified, only specialized, extended, or profiled.")
for i, (h, b, c) in enumerate([
    ("SPECIALIZE — add", "Type values, not parallel classes.\nAttribute overrides: add, remove,\nretighten. Subclasses under the\nconstruct roots.  — Annex B", BLUE),
    ("EXTEND — carry", "Open attribute register: dialect keys\non any construct, zero model change.\nRaw, variables, unknown directives\nall carried losslessly.  — Annex C", ACCENT),
    ("TAILOR — subtract", "Profiles: ~20 lines of YAML declare\nyour subset. Narrowing only — the\nmachine rejects any widening. A\nprofile document is always valid.", INK),
]):
    card(s, 0.7 + i * 2.95, 1.7, 2.75, 2.35, h, b, head_color=c)
box(s, 0.7, 4.3, 8.6, 0.5,
    "legalexchange/1.0 ships as the worked profile: excluded constructs, narrowed cardinalities, closed key vocabulary — validated in CI.",
    size=11, color=MUTE)
pageno(s)

# ---------------- 7 · Executable conformance ----------------
s = slide(CONTENT_LAYOUT)
set_title(s, "A standard you can run (Annex D)",
          "Conformance classes map one-to-one to gates that run on every change — the suite is the standard's proof.")
y = 1.55
for lab, body in [
    ("MODEL INTEGRITY", "Names resolve, every attribute carries visibility and a definition — checked on the definition modules."),
    ("DIAGRAM PARITY", "Every diagram view is provably the model: includes and associations in closure, no class bodies under views."),
    ("TWIN INSTANCE SUITES", "Identical XML and YAML reference instances exercise every construct — validated against schema and parsed model."),
    ("PROFILE CONFORMANCE", "Narrowing-only enforced mechanically; profile instances carry no excluded construct and permitted keys only."),
]:
    row(s, y, lab, body)
    y += 0.88
box(s, 0.7, 5.05, 8.6, 0.4,
    "“A construct without a serializable instance is unfinished” — enforced, not aspired to.",
    size=12, color=INK, bold=True)
pageno(s)

# ---------------- 8 · Production-proven ----------------
s = slide(CONTENT_LAYOUT)
set_title(s, "Extracted from production, not written for a committee",
          "The model underlies real standards authoring and publishing across 30+ document flavours today.")
for i, (h, b) in enumerate([
    ("LIVE ECOSYSTEM", "relaton-models and standoc-models consume it in lockstep; grammars compile and pass regeneration parity in their CI."),
    ("REGISTRY INTEGRATION", "Localized strings identify ISO 24229 spelling systems; romanization schemes by system code."),
    ("PUBLIC ARTIFACTS", "Model atlas with full definitions and hyperlinked types; grammar packages site; open repositories throughout."),
    ("GOVERNED", "Versioning policy with deprecation cycle, standardization snapshots, and a public register-key registry."),
]):
    card(s, 0.7 + i * 2.2, 1.7, 2.05, 2.0, h, b)
rect(s, 0.7, 3.95, 8.6, 0.7, fill=PANEL)
box(s, 0.9, 4.08, 8.2, 0.5,
    "Six years in production, then standardized — the reverse of the usual order.",
    size=13, color=INK, bold=True)
pageno(s)

# ---------------- 9 · Why TC 154 ----------------
s = slide(CONTENT_LAYOUT)
set_title(s, "TC 154's portfolio already standardizes the document's elements",
          "Processes, data elements and documents in commerce — this is the layer the portfolio is missing.")
for i, (h, b) in enumerate([
    ("ISO 8601", "when — time and date"),
    ("ISO 6523", "who — organization identifiers"),
    ("E-BUSINESS FRAMEWORKS", "structured trade data — rich text punted on"),
    ("ISO 690 · 5127 (TC 46)", "what is cited — consumed, not duplicated"),
]):
    card(s, 0.7 + i * 2.2, 1.7, 2.05, 1.35, h, b)
rect(s, 0.7, 3.3, 8.6, 1.1, fill=PAPER, line=ACCENT)
box(s, 0.95, 3.42, 8.1, 0.9,
    "The missing piece: WHAT THE DOCUMENT SAYS — the rich-text layer.\nOrders, invoices, contracts, specifications: every business document is lightweight rich text.",
    size=13, color=INK, bold=True)
box(s, 0.7, 4.6, 8.6, 0.6,
    "Time, parties, trade data, citations — standardized. The document carrying them: CC 36010 completes the stack.",
    size=12, color=BODY)
pageno(s)

# ---------------- 10 · The ask (template closing layout) ----------------
s = slide(CLOSING_LAYOUT)
box(s, 1.1, 0.75, 7.8, 0.5, "PUBLISHED AND READY", size=11, color=ACCENT,
    bold=True, align=PP_ALIGN.CENTER)
box(s, 0.7, 1.25, 8.6, 1.1,
    "Adopt CC 36010 as the base text\nfor a fast-track DIS",
    size=26, color=INK, bold=True, align=PP_ALIGN.CENTER)
bar(s, 4.4, 2.5, w=1.2)
for i, (h, b) in enumerate([
    ("PUBLISHED", "CC 36010 published by CalConnect, 2026-08-31 — stable base text with machine-checked annexes."),
    ("STAGED", "The ISO copy stands at DIS stage 40.00 — identical sources, two flavours, fast-track ready."),
    ("MAINTAINED", "Open repositories, executable test suite, versioning policy — the maintenance vehicle exists today."),
]):
    card(s, 0.7 + i * 2.95, 2.85, 2.75, 1.7, h, b)
box(s, 0.7, 4.75, 8.6, 0.6,
    "calconnect.github.io/cc-lightweight-doc · metanorma.github.io/basicdoc-models\nCalConnect TC VCARD · Ronald Tse",
    size=10.5, color=MUTE, align=PP_ALIGN.CENTER)
box(s, 9.3, 5.32, 0.3, 0.2, "10", size=6.5, color=MUTE)

out = ROOT / "presentations" / "20260831-calconnect-iso-tc-154-cc36010.pptx"
prs.save(str(out))
print(f"saved: {out}")
